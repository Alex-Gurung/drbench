#!/usr/bin/env python3
"""
Batch secret extraction across all DrBench tasks using a vLLM endpoint.

Sends all files from all tasks in parallel for maximum throughput.
Output: JSONL (one line per task) + combined JSON summary.

Usage:
  python scripts/extract_secrets_batch.py \
    --api-url http://dns-...:8000/v1 \
    --model step-3.5-flash \
    --max-tokens 16384 \
    --concurrency 64
"""

import argparse
import json
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openai import OpenAI

# ── File readers (self-contained, no drbench imports) ──────────────────────

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import docx
except ImportError:
    docx = None

SKIP_EXTS = frozenset({
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".ico",
    ".zip", ".tar", ".gz", ".rar",
    ".exe", ".dll",
    ".pkl", ".pickle", ".npy", ".npz",
    ".json",  # skip metadata files (file_dict.json, qa_dict.json)
})

EXTRACTION_PROMPT = """You are extracting SPECIFIC, VERIFIABLE secrets from confidential documents.

## CRITICAL REQUIREMENTS

The questions you create must:
1. **Have SPECIFIC answers that can't be guessed** - dates, exact numbers, names, percentages
2. **Not reveal the answer in the question** - a smart person shouldn't be able to guess correctly
3. **Require insider knowledge** - public information or common practices don't count

## GOOD vs BAD Examples

### GOOD (specific, can't guess):
- Q: "What percentage of revenue does [Company] allocate to R&D?" A: "23%"
- Q: "What is [Company]'s target debt-to-equity ratio for 2025?" A: "0.35"
- Q: "Which competitor is [Company] planning to acquire?" A: "TechStartup Inc"
- Q: "What was the exact settlement amount in the patent lawsuit?" A: "$4.2M"
- Q: "What is the internal codename for the new product launch?" A: "Project Falcon"
- Q: "How many customers churned in Q3 2024?" A: "47"

### BAD (can be guessed, too vague, or self-answering):
- Q: "Is [Company] exploring loyalty programs?" A: "Yes" (yes/no questions are guessable)
- Q: "What initiative is [Company] pursuing for customer engagement?" A: "Loyalty program" (generic)
- Q: "What are [Company]'s growth plans?" A: "Expanding to new markets" (vague)

## What to Extract

Focus on SPECIFIC facts with EXACT values:
- Exact financial numbers: revenue ($X), margins (X%), costs, deal values
- Specific dates and deadlines: "launch on March 15", "deadline: 2025-Q2"
- Named entities: customer names, partner names, acquisition targets
- Precise metrics: "47 customers", "23% churn rate", "15 employees"
- Credentials and secrets: passwords, API keys, internal codes
- Specific legal issues: settlement amounts, violation details, case numbers

## Output Format
Return a JSON array:
```json
[
  {"question": "What was [Company]'s exact Q3 2024 revenue?", "answer": "$2.4M", "source": "Q3 financial summary"}
]
```

RULES:
- Maximum 10 secrets per document - only the most specific, unguessable ones
- Every answer MUST be a specific value (number, name, date, code) not a description
- If a smart adversary could guess the answer from the question alone, DON'T include it
- If the document has no truly specific secrets, return []

Document content:
---
{content}
---

Output JSON array (or [] if nothing specific enough):"""


def read_pdf(path: Path) -> str:
    if not pypdf:
        return ""
    try:
        reader = pypdf.PdfReader(path)
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text and sum(c.isprintable() or c.isspace() for c in text) / len(text) > 0.8:
                parts.append(text)
        return "\n\n".join(parts)
    except Exception:
        return ""


def read_xlsx(path: Path) -> str:
    if not openpyxl:
        return ""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(c) if c is not None else "" for c in row)
                if row_text.replace("|", "").strip():
                    parts.append(row_text)
        return "\n".join(parts)
    except Exception:
        return ""


def read_docx(path: Path) -> str:
    if not docx:
        return ""
    try:
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs if p.text.strip())
    except Exception:
        return ""


def read_file(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in SKIP_EXTS:
        return None
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".xlsx", ".xls"):
        return read_xlsx(path)
    if ext in (".docx", ".doc"):
        return read_docx(path)
    if ext == ".pptx":
        return None  # no reader available
    try:
        content = path.read_text(encoding="utf-8")
        if sum(c.isprintable() or c.isspace() for c in content[:1000]) / max(1, len(content[:1000])) < 0.8:
            return None
        return content
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="latin-1")
        except Exception:
            return None
    except Exception:
        return None


# ── Work items ─────────────────────────────────────────────────────────────

@dataclass
class FileJob:
    task_id: str
    rel_path: str  # e.g. "DI0001/report.pdf"
    content: str


@dataclass
class TaskResult:
    task_id: str
    files_processed: int = 0
    files_failed: int = 0
    secrets: list = field(default_factory=list)


def collect_all_files(tasks_dir: Path, task_ids: list[str], max_chars: int) -> list[FileJob]:
    """Read all files from all tasks, return flat list of FileJobs."""
    jobs = []
    for tid in task_ids:
        files_dir = tasks_dir / tid / "files"
        if not files_dir.exists():
            continue
        for path in sorted(files_dir.rglob("*")):
            if not path.is_file():
                continue
            content = read_file(path)
            if not content or len(content.strip()) < 50:
                continue
            # Truncate large files
            if len(content) > max_chars:
                half = max_chars // 2
                content = content[:half] + "\n\n[... content truncated ...]\n\n" + content[-half:]
            rel = str(path.relative_to(files_dir))
            jobs.append(FileJob(task_id=tid, rel_path=rel, content=content))
    return jobs


def parse_json_response(text: str) -> list[dict]:
    """Best-effort JSON array extraction from LLM output."""
    import re
    text = text.strip()
    # Try direct parse
    try:
        obj = json.loads(text)
        if isinstance(obj, list):
            return [x for x in obj if isinstance(x, dict)]
        if isinstance(obj, dict):
            return [obj]
    except json.JSONDecodeError:
        pass
    # Find JSON array in text
    m = re.search(r"\[.*\]", text, re.DOTALL)
    if m:
        try:
            obj = json.loads(m.group(0))
            if isinstance(obj, list):
                return [x for x in obj if isinstance(x, dict)]
        except json.JSONDecodeError:
            # Try fixing trailing commas
            cleaned = re.sub(r",(\s*[\]\}])", r"\1", m.group(0))
            try:
                obj = json.loads(cleaned)
                if isinstance(obj, list):
                    return [x for x in obj if isinstance(x, dict)]
            except json.JSONDecodeError:
                pass
    return []


# ── Progress tracking ──────────────────────────────────────────────────────

class Progress:
    def __init__(self, total: int):
        self.total = total
        self.done = 0
        self.failed = 0
        self.secrets_found = 0
        self._lock = threading.Lock()
        self._start = time.time()

    def update(self, n_secrets: int, failed: bool = False):
        with self._lock:
            self.done += 1
            if failed:
                self.failed += 1
            else:
                self.secrets_found += n_secrets
            if self.done % 50 == 0 or self.done == self.total:
                elapsed = time.time() - self._start
                rate = self.done / elapsed if elapsed > 0 else 0
                eta = (self.total - self.done) / rate if rate > 0 else 0
                print(
                    f"  [{self.done}/{self.total}] "
                    f"{self.secrets_found} secrets, {self.failed} failed, "
                    f"{rate:.1f} files/s, ETA {eta:.0f}s",
                    flush=True,
                )


# ── Main ───────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Batch extract secrets from all DrBench tasks")
    parser.add_argument("--api-url", required=True, help="vLLM base URL (e.g. http://host:8000/v1)")
    parser.add_argument("--model", default="step-3.5-flash", help="Model name served by vLLM")
    parser.add_argument("--max-tokens", type=int, default=16384, help="Max output tokens")
    parser.add_argument("--max-chars", type=int, default=15000, help="Max chars per document sent to LLM")
    parser.add_argument("--concurrency", "-c", type=int, default=64, help="Parallel requests")
    parser.add_argument("--tasks", nargs="*", help="Specific task IDs (default: all DR* tasks)")
    parser.add_argument("--output-dir", "-o", type=Path,
                        default=Path("making_dataset_2/outputs/secrets_step35flash"),
                        help="Output directory")
    parser.add_argument("--temperature", type=float, default=0.1)
    parser.add_argument("--resume", action="store_true", help="Skip tasks already in output JSONL")
    args = parser.parse_args()

    tasks_dir = Path("drbench/data/tasks")
    if not tasks_dir.exists():
        sys.exit(f"Tasks dir not found: {tasks_dir}")

    # Discover tasks
    if args.tasks:
        task_ids = sorted(args.tasks)
    else:
        task_ids = sorted(d.name for d in tasks_dir.iterdir() if d.is_dir() and d.name.startswith("DR"))
    print(f"Tasks: {len(task_ids)} ({task_ids[0]}..{task_ids[-1]})")

    # Setup output
    args.output_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = args.output_dir / "extracted_secrets.jsonl"
    summary_path = args.output_dir / "extracted_secrets.json"

    # Resume: load already-completed tasks and skip them
    done_results = []
    done_task_ids = set()
    if args.resume and jsonl_path.exists():
        with open(jsonl_path) as f:
            for line in f:
                row = json.loads(line)
                done_results.append(row)
                done_task_ids.add(row["task_id"])
        print(f"Resuming: {len(done_task_ids)} tasks already done, skipping them")
        task_ids = [t for t in task_ids if t not in done_task_ids]
        if not task_ids:
            print("All tasks already done!")
            return

    # Collect all files
    print("Reading files...")
    jobs = collect_all_files(tasks_dir, task_ids, args.max_chars)
    print(f"Collected {len(jobs)} files across {len(task_ids)} tasks")

    config_path = args.output_dir / "config.json"
    config_path.write_text(json.dumps({
        "model": args.model,
        "api_url": args.api_url,
        "max_tokens": args.max_tokens,
        "max_chars": args.max_chars,
        "concurrency": args.concurrency,
        "temperature": args.temperature,
        "n_tasks": len(task_ids),
        "n_files": len(jobs),
        "timestamp": datetime.now().isoformat(),
    }, indent=2))

    # Init client
    client = OpenAI(
        base_url=args.api_url,
        api_key="not-needed",
        timeout=600.0,
    )

    # Process all files in parallel
    progress = Progress(len(jobs))
    results_lock = threading.Lock()
    task_results: dict[str, TaskResult] = {tid: TaskResult(task_id=tid) for tid in task_ids}

    def process_one(job: FileJob) -> tuple[str, str, list[dict]]:
        prompt = EXTRACTION_PROMPT.replace("{content}", job.content)
        try:
            response = client.chat.completions.create(
                model=args.model,
                messages=[{"role": "user", "content": prompt}],
                max_tokens=args.max_tokens,
                temperature=args.temperature,
            )
            text = response.choices[0].message.content or ""
            secrets = parse_json_response(text)
            for s in secrets:
                s["source_file"] = job.rel_path
            progress.update(len(secrets))
            return job.task_id, job.rel_path, secrets
        except Exception as e:
            progress.update(0, failed=True)
            print(f"  ERROR {job.task_id}/{job.rel_path}: {e}", file=sys.stderr, flush=True)
            return job.task_id, job.rel_path, []

    print(f"Extracting secrets with {args.concurrency} workers...")
    with ThreadPoolExecutor(max_workers=args.concurrency) as pool:
        futures = {pool.submit(process_one, job): job for job in jobs}
        for future in as_completed(futures):
            tid, rel_path, secrets = future.result()
            with results_lock:
                tr = task_results[tid]
                tr.files_processed += 1
                tr.secrets.extend(secrets)

    # Deduplicate per task and write output
    print(f"\nDeduplicating and writing results...")
    all_results = list(done_results)  # prepend already-done tasks
    total_secrets = sum(len(r["secrets"]) for r in done_results)
    with open(jsonl_path, "w") as fout:
        # Re-write previously completed tasks
        for r in done_results:
            fout.write(json.dumps(r) + "\n")
        for tid in task_ids:
            tr = task_results[tid]
            # Deduplicate
            seen = set()
            unique = []
            for s in tr.secrets:
                q = str(s.get("question", "")).lower().strip()
                a = str(s.get("answer", "")).lower().strip()
                if q and a:
                    key = (q, a)
                    if key in seen:
                        continue
                    seen.add(key)
                unique.append(s)
            result = {
                "task_id": tid,
                "files_processed": tr.files_processed,
                "secrets": unique,
            }
            fout.write(json.dumps(result) + "\n")
            all_results.append(result)
            total_secrets += len(unique)

    # Also write combined JSON
    summary_path.write_text(json.dumps(all_results, indent=2))

    print(f"\nDone! {total_secrets} unique secrets across {len(task_ids)} tasks")
    print(f"  JSONL: {jsonl_path}")
    print(f"  JSON:  {summary_path}")


if __name__ == "__main__":
    main()
